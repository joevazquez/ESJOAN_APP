"""Prueba de humo de la interfaz con Playwright + mock de Supabase (no toca ninguna base real).
Uso (desde la raíz del proyecto):
    python3 -m http.server 8765 &
    pip install playwright openpyxl && playwright install chromium
    python3 tests/ui_smoke.py
Las capturas quedan en tests/_out/ (ignorado por git).
"""
import json, sys, re, io, os
from pathlib import Path
RAIZ = Path(__file__).resolve().parents[1]
SALIDA = RAIZ / 'tests' / '_out'; SALIDA.mkdir(exist_ok=True)
from playwright.sync_api import sync_playwright
BASE='http://localhost:8765/'
FIX=open(RAIZ/'tests'/'fixtures.json',encoding='utf8').read()
MOCK=open(RAIZ/'tests'/'mock-supabase.js',encoding='utf8').read()
CONFIG="export const SUPABASE_URL='https://test.supabase.co';export const SUPABASE_ANON_KEY='anon-test';export const BUCKET='expedientes';export const GASTOS_NOTARIALES_DEFECTO=0.05;"
errores=[]
def nuevo(browser, w=1280, h=860):
    ctx=browser.new_context(viewport={'width':w,'height':h}, accept_downloads=True)
    ctx.add_init_script("window.__FIX="+FIX+";")
    ctx.route("**/js/config.js", lambda r: r.fulfill(body=CONFIG, content_type='text/javascript'))
    ctx.route("**/js/vendor/supabase.js", lambda r: r.fulfill(body=MOCK, content_type='text/javascript'))
    p=ctx.new_page()
    p.on('pageerror', lambda e: errores.append('PAGEERROR '+str(e)))
    p.on('console', lambda m: errores.append('CONSOLE '+m.type+' '+m.text) if m.type in ('error','warning') else None)
    return ctx,p
def entrar(p, user):
    p.goto(BASE); p.evaluate("u=>localStorage.setItem('mockUser', JSON.stringify(u))", user); p.goto(BASE+'#/'); p.reload(); p.wait_for_selector('.side, .acceso', timeout=8000); p.wait_for_timeout(500)
def shot(p, name, full=False): p.screenshot(path=str(SALIDA/f'{name}.png'), full_page=full)
PDF=b'%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF'
with sync_playwright() as pw:
    b=pw.chromium.launch()
    # ---------- login (sin sesión)
    ctx,p=nuevo(b); p.goto(BASE); p.wait_for_selector('.acceso'); p.wait_for_timeout(400); shot(p,'01_login'); 
    p.goto(BASE+'#/registro'); p.wait_for_selector('#nombre'); shot(p,'02_registro')
    ctx.close()
    ctx,p=nuevo(b,390,800); p.goto(BASE); p.wait_for_selector('.acceso'); p.wait_for_timeout(300); shot(p,'03_login_movil'); ctx.close()

    # ---------- cliente
    ctx,p=nuevo(b); entrar(p,'cliente'); shot(p,'10_cliente_inicio')
    p.click('.exp-card'); p.wait_for_selector('.cumbrera svg'); p.wait_for_timeout(300); shot(p,'11_cliente_resumen', True)
    p.click('text=Documentos'); p.wait_for_selector('.doc-req'); shot(p,'12_cliente_docs', True)
    # subir un PDF válido a todos los obligatorios
    n=p.locator('input[type=file]').count(); print('inputs archivo',n)
    tipos=[t for t in json.loads(FIX)['tipos'] if t['obligatorio'] and (not t['aplica_tipos'] or 'adquisicion' in t['aplica_tipos'])]
    for t in tipos:
        p.set_input_files('#f_'+t['clave'], files=[{'name':t['clave']+'.pdf','mimeType':'application/pdf','buffer':PDF}]); p.wait_for_timeout(250)
    p.wait_for_selector('.cumbrera'); shot(p,'13_cliente_docs_completos', True)
    print('estatus tras subir:', p.inner_text('.estado-actual .lead'))
    # archivo no-PDF debe rechazarse
    p.set_input_files('#f_predial', files=[{'name':'malo.pdf','mimeType':'application/pdf','buffer':b'hola no soy pdf'}]); p.wait_for_timeout(300); shot(p,'14_cliente_pdf_falso')
    print('toast:', p.locator('.toast').all_inner_texts())

    # ---------- broker
    entrar(p,'broker'); shot(p,'20_broker_bandeja')
    p.click('a.fila-link >> nth=0'); p.wait_for_selector('.cumbrera svg'); shot(p,'21_broker_resumen',True)
    p.click('text=Documentos'); p.wait_for_selector('.doc-req'); shot(p,'22_broker_docs',True)
    while p.locator('button:has-text("Aprobar")').count():
        p.locator('button:has-text("Aprobar")').first.click(); p.wait_for_timeout(200)
    p.wait_for_timeout(400); p.click('text=Resumen'); p.wait_for_selector('.cumbrera'); print('estatus tras aprobar:', p.inner_text('.estado-actual .lead'))
    p.fill('#av_v','Valuación Integral SA'); p.fill('#av_f','2026-10-05'); p.click('button:has-text("Asignar avalúo")'); p.wait_for_timeout(500); print('->', p.inner_text('.estado-actual .lead'))
    p.fill('#fi_f','2026-10-20T11:00'); p.fill('#fi_n','Notaría 45 CDMX'); p.click('button:has-text("Asignar firma")'); p.wait_for_timeout(500); print('->', p.inner_text('.estado-actual .lead'))
    shot(p,'23_broker_firma',True)
    # simulador
    p.click('a.tab:has-text("Simulador")'); p.wait_for_selector('.sim-layout'); p.wait_for_selector('.oferta'); p.wait_for_timeout(400); shot(p,'24_simulador',True)
    print('ofertas:', p.locator('.oferta').count())
    p.locator('.oferta').first.click(); p.wait_for_selector('.modal'); p.wait_for_timeout(300); shot(p,'25_sim_detalle')
    p.locator('.modal').evaluate("m=>m.scrollTo?m.scrollTo(0,9999):0"); 
    p.click('button:has-text("Guardar en el expediente")'); p.wait_for_selector('#g_n'); p.check('#g_c'); p.fill('#g_no','Mejor equilibrio entre mensualidad y costo total.'); p.click('.modal-foot button:has-text("Guardar")'); p.wait_for_timeout(500)
    p.click('button:has-text("Guardadas")'); p.wait_for_selector('.hoja .oferta'); p.click('button:has-text("Marcar como recomendada")'); p.wait_for_timeout(400); shot(p,'26_sim_guardadas',True)
    # plazos múltiples
    p.click('button:has-text("Comparar bancos")'); p.wait_for_selector('.oferta'); 
    p.locator('label.chip:has-text("15 años")').click(); p.locator('label.chip:has-text("25 años")').click(); p.wait_for_timeout(600); print('ofertas 3 plazos:', p.locator('.oferta').count()); shot(p,'27_sim_multiplazo',True)

    # ---------- cliente ve opciones compartidas
    entrar(p,'cliente'); p.goto(BASE+'#/expediente/11111111-1111-4111-8111-111111111111/simulador'); p.wait_for_selector('.hoja'); p.wait_for_timeout(400); shot(p,'28_cliente_opciones',True)
    # ---------- admin
    entrar(p,'admin'); shot(p,'30_admin_bandeja')
    p.goto(BASE+'#/admin/usuarios'); p.wait_for_selector('.tabla'); shot(p,'31_admin_usuarios')
    p.goto(BASE+'#/admin/bancos'); p.wait_for_selector('.tabla'); shot(p,'32_admin_bancos')
    p.goto(BASE+'#/admin/bancos/banorte'); p.wait_for_selector('.tabla input'); p.wait_for_timeout(300); shot(p,'33_admin_banco',True)
    p.goto(BASE+'#/admin/documentos'); p.wait_for_selector('.tabla'); shot(p,'34_admin_docs')
    p.goto(BASE+'#/perfil'); p.wait_for_selector('.hoja'); shot(p,'35_perfil')
    ctx.close()
    b.close()


import openpyxl
with sync_playwright() as pw:
    b=pw.chromium.launch(); ctx,p=nuevo(b); entrar(p,'admin')
    p.goto(BASE+'#/admin/bancos/hsbc'); p.wait_for_selector('.tabla input'); p.wait_for_timeout(400)
    print('filas hsbc en grid:', p.inner_text('.hoja-body .muted.small, .row .muted.small')[:60])
    # --- descargar Excel
    with p.expect_download() as d: p.click('button:has-text("Descargar Excel")')
    dl=d.value; path=str(SALIDA/dl.suggested_filename); dl.save_as(path); print('descarga:', dl.suggested_filename)
    wb=openpyxl.load_workbook(path); ws=wb.active; print('xlsx filas', ws.max_row, 'cols', ws.max_column, [c.value for c in ws[1]][:8], [c.value for c in ws[2]][:9])
    # --- modificar tasas (+1 pp) y reimportar
    for r in range(2, ws.max_row+1): ws.cell(r, 7).value = round(float(ws.cell(r,7).value)+1, 4)
    wb.save(str(SALIDA/'tarifas_mod.xlsx'))
    p.set_input_files('#imp_arch', str(SALIDA/'tarifas_mod.xlsx')); p.wait_for_selector('.modal'); p.click('.modal-foot button:has-text("Cargar en pantalla")'); p.wait_for_timeout(600)
    print('toast import:', p.locator('.toast').all_inner_texts()[-1:])
    p.click('button:has-text("Guardar tarifas")'); p.wait_for_timeout(600)
    db=p.evaluate("()=>{const b=window.__DB.bancos.find(x=>x.clave==='hsbc');const t=window.__DB.bancos_tasas.filter(x=>x.banco_id===b.id);return {n:t.length, tasa:t[0].tasa, plazo:[t[0].plazo_min,t[0].plazo_max], aforo:t[0].aforo_max}}")
    print('DB tras importar xlsx:', db)
    # --- pegar TSV con % en puntos porcentuales
    p.goto(BASE+'#/admin/bancos/afirme'); p.wait_for_selector('.tabla input'); p.wait_for_timeout(300)
    p.click('button:has-text("Pegar desde Excel")')
    tsv="producto\tdestino\tprograma\tesquema\tperfil\ttasa\tplazo_min\tplazo_max\taforo_max\tcomision_apertura\n"+"Hipoteca fija\tadquisicion\t*\tfijo\tTasa promo\t9.75\t5\t20\t90\t1.25\nHipoteca fija\tadquisicion\t*\tfijo\tTasa base\t10.4\t5\t20\t90\t0.75\n"
    p.fill('textarea[aria-label="Datos pegados"]', tsv); p.click('.modal-foot button:has-text("Cargar en pantalla")'); p.wait_for_timeout(300)
    p.click('.modal-foot button:has-text("Cargar en pantalla")') if p.locator('.modal').count() else None
    p.wait_for_timeout(300); shot(p,'40_pegado')
    p.click('button:has-text("Guardar tarifas")'); p.wait_for_timeout(500)
    print('afirme:', p.evaluate("()=>{const b=window.__DB.bancos.find(x=>x.clave==='afirme');return window.__DB.bancos_tasas.filter(x=>x.banco_id===b.id).map(t=>[t.perfil,t.tasa,t.aforo_max,t.comision_apertura])}"))
    # tasa mal escalada debe bloquearse
    p.goto(BASE+'#/admin/bancos/afirme'); p.wait_for_selector('.tabla input'); p.wait_for_timeout(300)
    p.fill('tbody tr:first-child input[aria-label="Tasa %"]','1050'); p.click('button:has-text("Guardar tarifas")'); p.wait_for_timeout(300); shot(p,'41_tasa_invalida')
    print('aviso:', p.locator('.banner.bad').all_inner_texts())
    # --- cliente móvil
    ctx.close(); ctx,p=nuevo(b,390,820); entrar(p,'cliente'); shot(p,'42_movil_inicio'); p.goto(BASE+'#/expediente/11111111-1111-4111-8111-111111111111'); p.wait_for_selector('.cumbrera'); p.wait_for_timeout(300); shot(p,'43_movil_expediente',True)
    p.click('button[aria-label="Abrir menú"]'); p.wait_for_timeout(400); shot(p,'44_movil_menu')
    print('\n'.join(sorted(set(errores))) or 'SIN ERRORES DE CONSOLA')
